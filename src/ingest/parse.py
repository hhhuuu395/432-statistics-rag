"""文档解析：pdf/docx/xlsx/pptx → 统一文本记录 [{text, page, source}]。"""
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from ..utils import config


def _clean(text):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    lines = [ln for ln in lines if not any(w in ln for w in config.WATERMARK_PATTERNS)]
    return '\n'.join(lines)


def parse_pdf(path):
    out = []
    for i, page in enumerate(fitz.open(str(path))):
        t = _clean(page.get_text())
        if t:
            out.append({'text': t, 'page': i + 1, 'source': str(path)})
    return out


def parse_docx(path):
    d = Document(str(path))
    paras = [p.text for p in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            paras.append(' | '.join(c.text for c in row.cells))
    t = _clean('\n'.join(paras))
    return [{'text': t, 'page': 1, 'source': str(path)}] if t else []


def parse_xlsx(path):
    wb = load_workbook(str(path), data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None]
            if cells:
                rows.append(' | '.join(cells))
    t = _clean('\n'.join(rows))
    return [{'text': t, 'page': 1, 'source': str(path)}] if t else []


def parse_pptx(path):
    prs = Presentation(str(path))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                texts.append(shape.text)
    t = _clean('\n'.join(texts))
    return [{'text': t, 'page': 1, 'source': str(path)}] if t else []


def parse_file(path):
    path = Path(path)
    handlers = {'pdf': parse_pdf, 'docx': parse_docx, 'xlsx': parse_xlsx, 'pptx': parse_pptx}
    h = handlers.get(path.suffix.lower().lstrip('.'))
    return h(path) if h else []


def parse_directory(dir_path):
    out = []
    for p in Path(dir_path).rglob('*'):
        if p.is_file():
            out.extend(parse_file(p))
    return out
